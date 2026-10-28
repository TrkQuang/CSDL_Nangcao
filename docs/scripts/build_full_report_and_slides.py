"""
Tự động biên dịch Báo cáo tổng kết hoàn chỉnh (.docx) và Slide thuyết trình (.pptx)
cho Đề tài 17 - Cơ sở dữ liệu nâng cao.
"""

from __future__ import annotations
import os, sys, re
from pathlib import Path
import docx
from docx import Document
from docx.shared import Inches, Pt, RGBColor, Cm
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.oxml import OxmlElement
from docx.oxml.ns import qn

import pptx
from pptx import Presentation
from pptx.util import Inches as PInches, Pt as PPt
from pptx.dml.color import RGBColor as PRGBColor
from pptx.enum.text import PP_ALIGN
from pptx.enum.shapes import MSO_SHAPE

ROOT = Path(__file__).resolve().parents[2]

def build_docx_report():
    print("=== Đang xây dựng Báo cáo tổng kết hoàn chỉnh (.docx) ===")
    doc = Document()
    
    # Cấu hình trang A4 chuẩn học thuật
    for section in doc.sections:
        section.page_width = Cm(21.0)
        section.page_height = Cm(29.7)
        section.top_margin = Cm(2.0)
        section.bottom_margin = Cm(2.0)
        section.left_margin = Cm(2.5)
        section.right_margin = Cm(2.0)
        
        # Header / Footer
        header = section.header
        hp = header.paragraphs[0]
        hp.text = "ĐỒ ÁN CƠ SỞ DỮ LIỆU NÂNG CAO — ĐỀ TÀI 17"
        hp.alignment = WD_ALIGN_PARAGRAPH.RIGHT
        hp.style.font.name = "Times New Roman"
        hp.style.font.size = Pt(8.5)
        hp.style.font.italic = True
        
        footer = section.footer
        fp = footer.paragraphs[0]
        fp.text = "Trường Đại học Sài Gòn — Khoa CNTT | "
        fp.alignment = WD_ALIGN_PARAGRAPH.RIGHT
        fp.style.font.name = "Times New Roman"
        fp.style.font.size = Pt(9)
        # add page number field
        fld = OxmlElement('w:fldSimple')
        fld.set(qn('w:instr'), 'PAGE')
        fp._p.append(fld)

    # Styles
    normal_style = doc.styles['Normal']
    normal_style.font.name = 'Times New Roman'
    normal_style.font.size = Pt(12)
    normal_style.paragraph_format.line_spacing = 1.2
    normal_style.paragraph_format.space_after = Pt(4)

    # --- TRANG BÌA ---
    p_uni = doc.add_paragraph()
    p_uni.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r_uni = p_uni.add_run("BỘ GIÁO DỤC VÀ ĐÀO TẠO\nTRƯỜNG ĐẠI HỌC SÀI GÒN\nKHOA CÔNG NGHỆ THÔNG TIN\n")
    r_uni.bold = True
    r_uni.font.size = Pt(13)
    p_uni.add_run("—" * 20 + "\n\n")

    p_title = doc.add_paragraph()
    p_title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r_report = p_title.add_run("BÁO CÁO TỔNG KẾT ĐỒ ÁN MÔN HỌC\nCƠ SỞ DỮ LIỆU NÂNG CAO\n\n")
    r_report.font.size = Pt(14)
    r_report.bold = True

    r_topic_label = p_title.add_run("ĐỀ TÀI SỐ 17:\n")
    r_topic_label.font.size = Pt(13)
    r_topic_label.bold = True
    r_topic_label.font.color.rgb = RGBColor(30, 64, 175)

    r_topic_name = p_title.add_run("DỮ LIỆU PHÂN CẤP VÀ DỮ LIỆU ĐỒ THỊ TRONG CSDL QUAN HỆ:\nỨNG DỤNG MẠNG XÃ HỘI HỌC TẬP\n\n")
    r_topic_name.font.size = Pt(16)
    r_topic_name.bold = True
    r_topic_name.font.color.rgb = RGBColor(17, 24, 39)

    doc.add_paragraph("\n" * 2)

    p_info = doc.add_paragraph()
    p_info.paragraph_format.left_indent = Inches(1.5)
    p_info.add_run("Giảng viên hướng dẫn: ").bold = True
    p_info.add_run("Bộ môn Hệ thống Thông tin\n")
    p_info.add_run("Nhóm sinh viên thực hiện: ").bold = True
    p_info.add_run("Nhóm Đề tài 17\n")
    p_info.add_run("Danh sách thành viên:\n").bold = True
    p_info.add_run("  1. Trần Quang (Nhóm trưởng)  — MSSV: 3124410292\n")
    p_info.add_run("  2. Thanh Hưng                 — Thành viên TV2\n")
    p_info.add_run("  3. Quang Hữu                  — Thành viên TV3\n")
    p_info.add_run("  4. Chương Võ                  — Thành viên TV4\n")
    p_info.add_run("  5. Như Quỳnh                  — MSSV: 3124410298\n\n")
    p_info.add_run("Hệ quản trị CSDL: ").bold = True
    p_info.add_run("PostgreSQL 16+ & Neo4j 5.x Community\n")
    p_info.add_run("Mã nguồn dự án: ").bold = True
    p_info.add_run("https://github.com/TrkQuang/CSDL_Nangcao.git\n")

    doc.add_paragraph("\n" * 3)
    p_date = doc.add_paragraph()
    p_date.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r_date = p_date.add_run("TP. HỒ CHÍ MINH — THÁNG 10/2026")
    r_date.bold = True
    r_date.font.size = Pt(12)

    doc.add_page_break()

    # Đọc nội dung Markdown chính đã tổng hợp
    md_file = ROOT / "report" / "de_tai_17_report_outline.md"
    content = md_file.read_text(encoding="utf-8")

    # Hàm parse markdown đơn giản và thêm vào Word
    lines = content.splitlines()
    in_code_block = False
    code_lines = []
    
    for line in lines:
        sline = line.strip()
        if sline.startswith("```"):
            if in_code_block:
                in_code_block = False
                # thêm khối code
                cp = doc.add_paragraph()
                cp.paragraph_format.left_indent = Inches(0.3)
                cp.paragraph_format.space_after = Pt(4)
                c_run = cp.add_run("\n".join(code_lines))
                c_run.font.name = "Consolas"
                c_run.font.size = Pt(9.5)
                c_run.font.color.rgb = RGBColor(30, 41, 59)
                code_lines = []
            else:
                in_code_block = True
                code_lines = []
            continue
            
        if in_code_block:
            code_lines.append(line)
            continue
            
        if not sline or sline == "---":
            continue
            
        if sline.startswith("# "):
            h = doc.add_heading(level=1)
            hr = h.add_run(sline[2:].strip())
            hr.font.name = "Times New Roman"
            hr.font.size = Pt(16)
            hr.bold = True
            hr.font.color.rgb = RGBColor(30, 58, 138)
            h.paragraph_format.space_before = Pt(12)
            h.paragraph_format.space_after = Pt(6)
        elif sline.startswith("## "):
            h = doc.add_heading(level=2)
            hr = h.add_run(sline[3:].strip())
            hr.font.name = "Times New Roman"
            hr.font.size = Pt(14)
            hr.bold = True
            hr.font.color.rgb = RGBColor(15, 23, 42)
            h.paragraph_format.space_before = Pt(10)
            h.paragraph_format.space_after = Pt(4)
        elif sline.startswith("### "):
            h = doc.add_heading(level=3)
            hr = h.add_run(sline[4:].strip())
            hr.font.name = "Times New Roman"
            hr.font.size = Pt(12.5)
            hr.bold = True
            hr.font.color.rgb = RGBColor(51, 65, 85)
            h.paragraph_format.space_before = Pt(8)
            h.paragraph_format.space_after = Pt(2)
        elif sline.startswith("|") and sline.endswith("|"):
            # Bảng
            continue
        elif sline.startswith("- ") or sline.startswith("* "):
            p = doc.add_paragraph(style='List Bullet')
            _add_formatted_text(p, sline[2:].strip())
        elif re.match(r"^\d+\.\s", sline):
            p = doc.add_paragraph(style='List Number')
            _add_formatted_text(p, re.sub(r"^\d+\.\s", "", sline).strip())
        else:
            p = doc.add_paragraph()
            _add_formatted_text(p, sline)

    out_docx = ROOT / "report" / "De_tai_17_Bao_cao_Tong_ket.docx"
    doc.save(str(out_docx))
    print(f"-> Đã xuất Báo cáo Word hoàn chỉnh: {out_docx} ({os.path.getsize(out_docx)} bytes)")

def _add_formatted_text(paragraph, text: str):
    tokens = re.split(r"(\*\*[^*]+\*\*|`[^`]+`|\$[^$]+\$)", text)
    for token in tokens:
        if token.startswith("**") and token.endswith("**"):
            r = paragraph.add_run(token[2:-2])
            r.bold = True
        elif token.startswith("`") and token.endswith("`"):
            r = paragraph.add_run(token[1:-1])
            r.font.name = "Consolas"
            r.font.size = Pt(10)
            r.font.color.rgb = RGBColor(180, 40, 40)
        elif token.startswith("$") and token.endswith("$"):
            r = paragraph.add_run(token[1:-1])
            r.font.italic = True
        else:
            # clean markdown links e.g. [text](url) -> text
            clean = re.sub(r"\[([^\]]+)\]\([^)]+\)", r"\1", token)
            paragraph.add_run(clean)

def build_pptx_slides():
    print("\n=== Đang xây dựng Slide PowerPoint thuyết trình (.pptx) ===")
    prs = Presentation()
    prs.slide_width = PInches(13.333)  # 16:9 Widescreen
    prs.slide_height = PInches(7.5)

    blank_layout = prs.slide_layouts[6]

    # Color Palette
    C_NAVY = PRGBColor(15, 23, 42)      # #0F172A
    C_BLUE = PRGBColor(37, 99, 235)     # #2563EB
    C_LIGHT_BG = PRGBColor(248, 250, 252) # #F8FAFC
    C_CARD_BG = PRGBColor(255, 255, 255)
    C_TEXT_DARK = PRGBColor(30, 41, 59)
    C_TEXT_MUTED = PRGBColor(100, 116, 139)
    C_ACCENT_GREEN = PRGBColor(16, 185, 129)
    C_ACCENT_ORANGE = PRGBColor(245, 158, 11)

    def add_header(slide, title: str, subtitle: str, speaker_tag: str):
        # Header banner
        header_box = slide.shapes.add_textbox(PInches(0.8), PInches(0.4), PInches(11.7), PInches(1.0))
        tf = header_box.text_frame
        tf.word_wrap = True
        tf.margin_left = tf.margin_top = tf.margin_right = tf.margin_bottom = 0
        
        p_tag = tf.paragraphs[0]
        p_tag.text = speaker_tag.upper()
        p_tag.font.size = PPt(11)
        p_tag.font.bold = True
        p_tag.font.color.rgb = C_BLUE
        
        p_title = tf.add_paragraph()
        p_title.text = title
        p_title.font.size = PPt(22)
        p_title.font.bold = True
        p_title.font.color.rgb = C_NAVY
        
        p_sub = tf.add_paragraph()
        p_sub.text = subtitle
        p_sub.font.size = PPt(13)
        p_sub.font.color.rgb = C_TEXT_MUTED

    # SLIDE 1: BÌA
    s1 = prs.slides.add_slide(blank_layout)
    bg1 = s1.shapes.add_shape(MSO_SHAPE.RECTANGLE, 0, 0, PInches(13.333), PInches(7.5))
    bg1.fill.solid()
    bg1.fill.fore_color.rgb = C_NAVY
    bg1.line.color.rgb = C_NAVY

    tb1 = s1.shapes.add_textbox(PInches(1.2), PInches(1.2), PInches(11.0), PInches(5.0))
    tf1 = tb1.text_frame
    tf1.word_wrap = True
    
    p = tf1.paragraphs[0]
    p.text = "BÁO CÁO ĐỒ ÁN MÔN HỌC — CƠ SỞ DỮ LIỆU NÂNG CAO"
    p.font.size = PPt(14)
    p.font.bold = True
    p.font.color.rgb = PRGBColor(147, 197, 253)

    p2 = tf1.add_paragraph()
    p2.text = "DỮ LIỆU PHÂN CẤP VÀ ĐỒ THỊ TRONG CSDL QUAN HỆ"
    p2.font.size = PPt(30)
    p2.font.bold = True
    p2.font.color.rgb = PRGBColor(255, 255, 255)

    p3 = tf1.add_paragraph()
    p3.text = "Ứng dụng Mạng Xã Hội Học Tập  |  Đề tài số 17"
    p3.font.size = PPt(20)
    p3.font.color.rgb = PRGBColor(226, 232, 240)

    p4 = tf1.add_paragraph()
    p4.text = "\n• Giảng viên hướng dẫn: Bộ môn Hệ thống Thông tin — Khoa CNTT, ĐH Sài Gòn"
    p4.font.size = PPt(14)
    p4.font.color.rgb = PRGBColor(203, 213, 225)

    p5 = tf1.add_paragraph()
    p5.text = "• Nhóm sinh viên thực hiện (5 thành viên):\n" \
              "   1. Trần Quang (Leader - 3124410292)  |  2. Thanh Hưng (TV2 - Cây & OOP)\n" \
              "   3. Quang Hữu (TV3 - SQL & Đệ quy)    |  4. Chương Võ (TV4 - Neo4j & Deadlock)\n" \
              "   5. Như Quỳnh (TV5 - Benchmark 3124410298)"
    p5.font.size = PPt(13.5)
    p5.font.color.rgb = PRGBColor(241, 245, 249)

    # SLIDE 2: ĐẶT VẤN ĐỀ & 3 CÂU HỎI NGHIÊN CỨU (TV1)
    s2 = prs.slides.add_slide(blank_layout)
    add_header(s2, "Đặt Vấn Đề & 3 Câu Hỏi Nghiên Cứu Cốt Lõi", 
               "Mạng xã hội học tập phát sinh dữ liệu phân cấp sâu và mạng lưới đồ thị phức tạp", "Phần 1 — Trần Quang (Leader)")
    
    # 3 Cards
    cards_data = [
        ("CÂU HỎI 1: MÔ HÌNH CÂY", 
         "Trong 4 mô hình (Adjacency, ltree, Nested Set, Closure Table), giải pháp nào tối ưu nhất cho từng thao tác (đọc cây con, tổ tiên, chèn, dời cành)? Chi phí đánh đổi về đĩa và tốc độ ghi ra sao?", 
         C_BLUE),
        ("CÂU HỎI 2: ĐỒ THỊ & WITH RECURSIVE", 
         "SQL đệ quy trên CSDL quan hệ đáp ứng các bài toán mạng xã hội (FoF, BFS đường ngắn nhất, chu trình) đến giới hạn nào? Khi nào cần chuyển đổi sang CSDL đồ thị chuyên dụng Neo4j?", 
         C_ACCENT_GREEN),
        ("CÂU HỎI 3: OOP, IMPEDANCE & TRANSACTION", 
         "Mô hình đối tượng ODL/OQL giải quyết bài toán Lệch pha trở kháng ra sao? Cơ chế cô lập giao dịch và kiểm soát bế tắc (Deadlock) vận hành thế nào dưới tải đồng thời cao?", 
         C_ACCENT_ORANGE),
    ]
    for i, (ctitle, cdesc, ccolor) in enumerate(cards_data):
        cshape = s2.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, PInches(0.8 + i*3.9), PInches(1.8), PInches(3.7), PInches(4.8))
        cshape.fill.solid()
        cshape.fill.fore_color.rgb = C_CARD_BG
        cshape.line.color.rgb = ccolor
        cshape.line.width = PPt(2)
        ctf = cshape.text_frame
        ctf.word_wrap = True
        ctf.margin_left = ctf.margin_right = ctf.margin_top = ctf.margin_bottom = PInches(0.25)
        
        cp1 = ctf.paragraphs[0]
        cp1.text = ctitle
        cp1.font.size = PPt(14)
        cp1.font.bold = True
        cp1.font.color.rgb = ccolor
        
        cp2 = ctf.add_paragraph()
        cp2.text = "\n" + cdesc
        cp2.font.size = PPt(12.5)
        cp2.font.color.rgb = C_TEXT_DARK

    # SLIDE 3: 10 QUY TẮC NGHIỆP VỤ & EER DIAGRAM (TV1)
    s3 = prs.slides.add_slide(blank_layout)
    add_header(s3, "Đặc Tả 10 Quy Tắc Nghiệp Vụ (BR1–BR10) & Mô Hình EER", 
               "Mô hình hóa chính xác ngữ nghĩa thế giới thực bằng hệ thống ràng buộc toàn vẹn", "Phần 1 — Trần Quang (Leader)")
    tb3 = s3.shapes.add_textbox(PInches(0.8), PInches(1.8), PInches(11.7), PInches(5.0))
    tf3 = tb3.text_frame
    tf3.word_wrap = True
    tf3.text = "• BR1 (Email duy nhất): UNIQUE(email) + Regex CHECK trên bảng users.\n" \
               "• BR2 (Bạn bè hai chiều vô hướng): Quy ước user_a < user_b; View logic friend_edge dùng UNION ALL hỗ trợ duyệt đồ thị.\n" \
               "• BR3 (Theo dõi một chiều có hướng): PRIMARY KEY (follower_id, followee_id), CHECK (follower <> followee).\n" \
               "• BR4 (Chủ nhóm & Quyền hạn): Chủ nhóm bắt buộc là thành viên; vai trò role IN ('owner', 'mod', 'member').\n" \
               "• BR5 (Quyền đăng bài): Chỉ thành viên chính thức mới được đăng bài vào nhóm học tập.\n" \
               "• BR6 (Toàn vẹn cây bình luận): Bình luận con bắt buộc cùng bài viết với bình luận cha; chặn chu trình cha-con.\n" \
               "• BR7 (Chuyên biệt hóa bài viết): Single Table Inheritance (post_thaoluan vs post_tailieu) — Rời nhau (d), Toàn phần (Total).\n" \
               "• BR8 (Cảm xúc duy nhất): PRIMARY KEY (post_id, user_id) trên bảng post_reaction.\n" \
               "• BR9 (Đồng bộ bộ đếm bạn bè): friend_count cập nhật nguyên tử trong cùng transaction kết bạn.\n" \
               "• BR10 (Bảo mật nhóm riêng tư): Quản lý quyền truy cập bài viết nhóm kín qua Row-Level Security."
    for p in tf3.paragraphs:
        p.font.size = PPt(13)
        p.font.color.rgb = C_TEXT_DARK

    # SLIDE 4: PHỤ THUỘC HÀM, CHUẨN HÓA BCNF & 4NF (TV1)
    s4 = prs.slides.add_slide(blank_layout)
    add_header(s4, "Lý Thuyết Phụ Thuộc Hàm, Chuẩn Hóa BCNF & Phân Rã 4NF", 
               "Chứng minh hình thức Lossless Join qua bảng ma trận Chase và phân rã triệt tiêu MVD", "Phần 1 — Trần Quang (Leader)")
    tb4 = s4.shapes.add_textbox(PInches(0.8), PInches(1.8), PInches(11.7), PInches(5.0))
    tf4 = tb4.text_frame
    tf4.word_wrap = True
    tf4.text = "• Lược đồ tổng hợp phẳng R vi phạm nghiêm trọng 3NF do các phụ thuộc bắc cầu qua thuộc tính không khóa:\n" \
               "   comment_id -> post_id -> group_id -> group_name  và  comment_id -> author_id -> author_name\n\n" \
               "• Phân rã theo thuật toán tổng hợp Bernstein thành 4 quan hệ con:\n" \
               "   R1(comment_id, post_id, author_id, parent_id, content, created_at) [Khóa: comment_id]\n" \
               "   R2(post_id, group_id) [Khóa: post_id]   |   R3(group_id, group_name, group_topic) [Khóa: group_id]\n" \
               "   R4(author_id, author_name, author_email) [Khóa: author_id]\n\n" \
               "• Kiểm tra BCNF: Trong mỗi quan hệ con, mọi định thức vế trái đều là siêu khóa => Đạt chuẩn BCNF 100%.\n" \
               "• Bảng ma trận Chase: Áp dụng các FD trong F_min, hàng 1 chuyển đổi thành toàn bộ ký hiệu 'a' => Xác nhận Lossless Join.\n" \
               "• Chuẩn 4NF: Tách thuộc tính đa trị độc lập (sở thích và môn học) thành user_interest và user_subject, triệt tiêu tích Descartes (tiết kiệm 80% dung lượng đĩa)."
    for p in tf4.paragraphs:
        p.font.size = PPt(13)
        p.font.color.rgb = C_TEXT_DARK

    # SLIDE 5: 4 MÔ HÌNH CÂY PHÂN CẤP (TV2 - HƯNG)
    s5 = prs.slides.add_slide(blank_layout)
    add_header(s5, "Cơ Sở Lý Thuyết & Cấu Trúc 4 Mô Hình Lưu Trữ Cây", 
               "So sánh đặc trưng kiến trúc giữa Adjacency List, ltree, Nested Set và Closure Table", "Phần 2 — Thanh Hưng (TV2)")
    # Thêm bảng tóm tắt 4 mô hình
    table_shape = s5.shapes.add_table(5, 5, PInches(0.8), PInches(1.8), PInches(11.7), PInches(4.8))
    table = table_shape.table
    headers = ["Mô hình cây", "Cột lưu trữ cốt lõi", "Ưu điểm nổi bật", "Nhược điểm lớn nhất", "Chỉ mục tối ưu"]
    for j, h in enumerate(headers):
        cell = table.cell(0, j)
        cell.text = h
        cell.fill.solid()
        cell.fill.fore_color.rgb = C_NAVY
        for p in cell.text_frame.paragraphs:
            p.font.bold = True
            p.font.size = PPt(12)
            p.font.color.rgb = PRGBColor(255, 255, 255)
    
    rows = [
        ["1. Adjacency List", "parent_id", "Chèn nút cực nhanh O(1), tối giản", "Đọc cây con phải dùng đệ quy CTE", "B-tree (post_id, parent_id)"],
        ["2. Path Enumeration", "path (ltree)", "Đọc cây con, tổ tiên siêu tốc bằng <@ / @>", "Chi phí sửa chuỗi path khi dời cành", "GiST (path)"],
        ["3. Nested Set", "lft, rgt", "Đọc cây con bằng 1 phép BETWEEN O(1)", "Chèn/xóa/dời cành phải re-index cả cây", "Composite (post_id, lft, rgt)"],
        ["4. Closure Table", "ancestor, descendant, depth", "Đọc cây con, tính độ sâu bằng JOIN cực nhanh", "Tốn dung lượng đĩa nhất O(N^2)", "Composite (ancestor, depth)"],
    ]
    for i, row in enumerate(rows):
        for j, val in enumerate(row):
            cell = table.cell(i+1, j)
            cell.text = val
            for p in cell.text_frame.paragraphs:
                p.font.size = PPt(11.5)
                p.font.color.rgb = C_TEXT_DARK

    # SLIDE 6: HIỆN THỰC BỘ 10 TRUY VẤN Q1–Q10 & CRUD (TV2 - HƯNG)
    s6 = prs.slides.add_slide(blank_layout)
    add_header(s6, "Bộ 10 Truy Vấn Chuẩn Q1–Q10 & Thao Tác Chuyển Nhánh An Toàn", 
               "Quy ước thống nhất Ground Truth và kiểm thử tự động xác nhận 4 mô hình đồng nhất", "Phần 2 — Thanh Hưng (TV2)")
    tb6 = s6.shapes.add_textbox(PInches(0.8), PInches(1.8), PInches(11.7), PInches(5.0))
    tf6 = tb6.text_frame
    tf6.word_wrap = True
    tf6.text = "• Quy ước Ground Truth chuẩn hóa:\n" \
               "   - Cây con / Hậu duệ (Q2, Q8): TÍNH CẢ chính nút gốc => subtree(1) = {1, 2, 3}, số lượng count(1) = 3.\n" \
               "   - Tổ tiên nghiêm ngặt (Q3, Q10): KHÔNG tính chính nút đang xét => ancestors(3) = {1, 2}.\n" \
               "   - Độ sâu (Q4): Nút gốc có depth = 1; Mức thụt lề (Q1): Nút gốc có level = 0.\n\n" \
               "• Thao tác CRUD và Chuyển nhánh (move_subtree) an toàn:\n" \
               "   - Kiểm tra chống chu trình: Không bao giờ cho phép chuyển nút A làm con của chính hậu duệ của A.\n" \
               "   - Adjacency List: Cập nhật parent_id sau khi kiểm tra CTE đệ quy.\n" \
               "   - Path Enumeration: Thay thế tiền tố path của toàn bộ cây con bằng hàm cắt chuỗi an toàn.\n" \
               "   - Nested Set: Thủ tục ns_move_subtree tính toán dịch chuyển khoảng cách lft/rgt.\n" \
               "   - Closure Table: Xóa các liên kết tổ tiên cũ và nhân bản các cặp liên kết tổ tiên mới.\n\n" \
               "• Kết quả kiểm thử tự động qua validate.ps1: 4/4 mô hình vượt qua kiểm tra queries=OK, crud=OK."
    for p in tf6.paragraphs:
        p.font.size = PPt(13)
        p.font.color.rgb = C_TEXT_DARK

    # SLIDE 7: SQL NÂNG CAO (Q1–Q8) & LẬP TRÌNH CSDL Q9–Q12 (TV3 - HỮU)
    s7 = prs.slides.add_slide(blank_layout)
    add_header(s7, "SQL Nâng Cao, Phép Chia Quan Hệ & Lập Trình Thủ Tục CSDL", 
               "Hiện thực phép chia 2 tầng NOT EXISTS, Window Functions và chống SQL Injection", "Phần 3 — Quang Hữu (TV3)")
    tb7 = s7.shapes.add_textbox(PInches(0.8), PInches(1.8), PInches(11.7), PInches(5.0))
    tf7 = tb7.text_frame
    tf7.word_wrap = True
    tf7.text = "• Phép chia quan hệ (Q3): Tìm sinh viên tham gia TẤT CẢ các nhóm học tập chủ đề 'CSDL':\n" \
               "   Dùng 2 tầng NOT EXISTS lồng nhau: 'Tìm SV sao cho KHÔNG TỒN TẠI nhóm CSDL nào mà SV đó KHÔNG tham gia'.\n\n" \
               "• Phân tích dữ liệu bằng Window Functions (Q5, Q6, Q7):\n" \
               "   - DENSE_RANK(): Xếp hạng người dùng tích cực nhất từng tháng mà không làm mất thông tin chi tiết.\n" \
               "   - LAG(): Tính khoảng cách thời gian giữa hai bài đăng liên tiếp của cùng một tác giả.\n\n" \
               "• Lập trình Cơ sở dữ liệu nâng cao (Q9–Q12):\n" \
               "   - Trigger Q9: Đảm bảo bình luận con và bình luận cha bắt buộc thuộc cùng một bài viết (BR6).\n" \
               "   - Procedure Q10 (proc_add_friend): Kết bạn nguyên tử, tự hoán vị LEAST/GREATEST và đồng bộ friend_count.\n" \
               "   - Procedure Q11 (proc_move_comment): Di chuyển nhánh cây an toàn, chặn chu trình đệ quy.\n" \
               "   - Dynamic SQL Q12: Dùng format('%L') chống SQL Injection tuyệt đối trước các chuỗi độc hại ' OR 1=1 --.\n" \
               "   - Embedded SQL: Module Python psycopg2 minh họa con trỏ Cursor và tham số hóa an toàn."
    for p in tf7.paragraphs:
        p.font.size = PPt(13)
        p.font.color.rgb = C_TEXT_DARK

    # SLIDE 8: SQL ĐỆ QUY TRÊN ĐỒ THỊ BẠN BÈ R1–R4 (TV3 - HỮU)
    s8 = prs.slides.add_slide(blank_layout)
    add_header(s8, "SQL Đệ Quy Trên Đồ Thị Bạn Bè Bằng WITH RECURSIVE", 
               "Duyệt đồ thị vô hướng 2 chiều qua View friend_edge và các bài toán mạng xã hội", "Phần 3 — Quang Hữu (TV3)")
    tb8 = s8.shapes.add_textbox(PInches(0.8), PInches(1.8), PInches(11.7), PInches(5.0))
    tf8 = tb8.text_frame
    tf8.word_wrap = True
    tf8.text = "• Cơ chế duyệt đồ thị 2 chiều: Xây dựng View friend_edge dùng UNION ALL hoán vị (user_a, user_b) và (user_b, user_a).\n\n" \
               "• R1 — Bạn của bạn (FoF độ sâu k = 2): Lấy bạn bè ở khoảng cách đúng 2 bước của An (user_id = 1), loại bỏ chính An và những người đã là bạn trực tiếp => Kết quả kỳ vọng: {Dung (ID 4), Em (ID 5)}.\n\n" \
               "• R2 — Đường đi ngắn nhất BFS: Thuật toán tìm kiếm theo chiều rộng duy trì mảng nút đã đi qua ARRAY[1, b] để chặn chu trình và dừng lại ngay khi chạm đỉnh đích => Đường đi ngắn nhất An -> Em: [1, 3, 5] (độ dài 2).\n\n" \
               "• R3 — Phát hiện chu trình vòng kín: Liệt kê các chu trình bạn bè khép kín độ dài 3 đến 5 => Phát hiện 3 chu trình: (1-2-3), (2-3-5-4) và (1-2-4-5-3).\n\n" \
               "• R4 — Thuật toán gợi ý kết bạn: Đếm số lượng bạn chung giữa người dùng mục tiêu với các người dùng chưa kết bạn, xếp hạng giảm dần theo độ ưu tiên đề xuất."
    for p in tf8.paragraphs:
        p.font.size = PPt(13)
        p.font.color.rgb = C_TEXT_DARK

    # SLIDE 9: CSDL HƯỚNG ĐỐI TƯỢNG ODL/OQL/JPQL (TV2 - HƯNG)
    s9 = prs.slides.add_slide(blank_layout)
    add_header(s9, "Mô Hình Hướng Đối Tượng ODL/OQL & Lệch Pha Trở Kháng", 
               "Thiết kế 6 lớp chuẩn ODMG 3.0 và phân tích Object-Relational Impedance Mismatch", "Phần 2 — Thanh Hưng (TV2)")
    tb9 = s9.shapes.add_textbox(PInches(0.8), PInches(1.8), PInches(11.7), PInches(5.0))
    tf9 = tb9.text_frame
    tf9.word_wrap = True
    tf9.text = "• Lược đồ ODL chuẩn ODMG 3.0 (schema.odl):\n" \
               "   - Định nghĩa 6 lớp: NguoiDung, NhomHocTap, BaiDang, BaiThaoLuan, BaiTaiLieu, BinhLuan.\n" \
               "   - Khai báo tập mở rộng (extent), quan hệ đối ngẫu 2 chiều (inverse), và các phương thức doSau(), soBanChung().\n\n" \
               "• Truy vấn OQL (oql_queries.md):\n" \
               "   - Điều hướng đối tượng tự nhiên bằng toán tử dấu chấm (.) không cần viết lệnh JOIN: p.tacGia.hoTen\n" \
               "   - Gọi phương thức trực tiếp trong mệnh đề lọc: SELECT c.noiDung FROM c IN dsBinhLuan WHERE c.doSau() >= 3;\n\n" \
               "• Phân tích Lệch pha trở kháng (Object-Relational Impedance Mismatch):\n" \
               "   1. Lệch pha kế thừa: OOP có kế thừa tự nhiên; RDBMS phải mô phỏng bằng Single Table kèm cột phân loại post_type.\n" \
               "   2. Lệch pha bản sắc (Identity): OOP phân biệt bằng địa chỉ bộ nhớ/OID; RDBMS phân biệt bằng giá trị khóa chính.\n" \
               "   3. Vấn đề N+1 truy vấn trong ORM (JPA/Hibernate): Phải dùng JOIN FETCH để tránh bắn thêm N câu truy vấn con."
    for p in tf9.paragraphs:
        p.font.size = PPt(13)
        p.font.color.rgb = C_TEXT_DARK

    # SLIDE 10: CSDL ĐỒ THỊ NEO4J & CYPHER VS POSTGRESQL (TV4 - CHƯƠNG)
    s10 = prs.slides.add_slide(blank_layout)
    add_header(s10, "CSDL Đồ Thị Chuyên Dụng Neo4j & Đối Chiếu Với PostgreSQL", 
               "So sánh toàn diện cú pháp và hiệu năng giữa Cypher và WITH RECURSIVE", "Phần 4 — Chương Võ (TV4)")
    
    # Text bên trái, hình bên phải
    tb10 = s10.shapes.add_textbox(PInches(0.8), PInches(1.8), PInches(6.2), PInches(5.0))
    tf10 = tb10.text_frame
    tf10.word_wrap = True
    tf10.text = "• Triển khai Neo4j Community: Nạp đồ thị 6 người dùng (:User) và các quan hệ [:FRIEND], [:FOLLOWS].\n\n" \
                "• So sánh cú pháp đường đi ngắn nhất:\n" \
                "   - Neo4j (Cypher): MATCH p = shortestPath((a:User {name:'An'})-[:FRIEND*..6]-(b:User {name:'Em'})) RETURN p;\n" \
                "   - PostgreSQL: CTE đệ quy BFS dài hơn 25 dòng kèm mảng kiểm tra vết.\n\n" \
                "• Đánh giá hiệu năng và kiến trúc:\n" \
                "   - Neo4j áp dụng Index-free Adjacency (duyệt con trỏ trực tiếp trong RAM) => Hiệu năng vượt trội khi độ sâu k >= 3.\n" \
                "   - PostgreSQL tối ưu cho các quan hệ phẳng, cây phân cấp và giao dịch ACID cốt lõi."
    for p in tf10.paragraphs:
        p.font.size = PPt(12.5)
        p.font.color.rgb = C_TEXT_DARK

    # Thêm ảnh chụp Neo4j
    img_neo4j = ROOT / "neo4j" / "screenshots" / "01_neo4j_full_graph.png"
    if img_neo4j.exists():
        s10.shapes.add_picture(str(img_neo4j), PInches(7.3), PInches(1.8), width=PInches(5.2))

    # SLIDE 11: ĐIỀU KHIỂN ĐỒNG THỜI & DEADLOCK T3 (TV4 - CHƯƠNG)
    s11 = prs.slides.add_slide(blank_layout)
    add_header(s11, "Điều Khiển Đồng Thời & Thực Nghiệm Kích Hoạt Deadlock (T3)", 
               "Tái hiện bế tắc khóa dòng giữa 2 phiên và ma trận kiểm thử các mức cô lập", "Phần 4 — Chương Võ (TV4)")
    tb11 = s11.shapes.add_textbox(PInches(0.8), PInches(1.8), PInches(11.7), PInches(5.0))
    tf11 = tb11.text_frame
    tf11.word_wrap = True
    tf11.text = "• Kịch bản thực nghiệm Deadlock giữa 2 phiên (Session A & Session B):\n" \
                "   - Session A xin khóa dòng User 1 FOR UPDATE, sau đó yêu cầu khóa User 2.\n" \
                "   - Session B xin khóa dòng User 2 FOR UPDATE, sau đó yêu cầu khóa User 1.\n" \
                "   - Kết quả thực nghiệm: PostgreSQL phát hiện chu trình chờ (Wait-For Graph) và hủy Session A với mã lỗi chính xác:\n" \
                "     ERROR: deadlock detected (SQLSTATE: 40P01) — Log nguyên bản lưu tại deadlock_t3_log.txt.\n\n" \
                "• Ma trận khảo sát 3 mức cô lập giao dịch trên PostgreSQL:\n" \
                "   - READ COMMITTED: Chặn Dirty Read; vẫn xảy ra Non-Repeatable Read và Phantom Read.\n" \
                "   - REPEATABLE READ: Áp dụng MVCC Snapshot Isolation => Chặn thành công cả Non-Repeatable Read và Phantom Read.\n" \
                "   - SERIALIZABLE: Thuật toán SSI theo dõi SIREAD locks, tự động ngắt phiên khi có xung đột tuần tự hóa (mã lỗi 40001)."
    for p in tf11.paragraphs:
        p.font.size = PPt(13)
        p.font.color.rgb = C_TEXT_DARK

    # SLIDE 12: THỬ TẢI ĐỒNG THỜI BẰNG PGBENCH (TV4 - CHƯƠNG)
    s12 = prs.slides.add_slide(blank_layout)
    add_header(s12, "Thử Tải Đồng Thời Bằng pgbench Trên Stored Procedure", 
               "Đo lường thông lượng TPS và độ trễ dưới các mức tải 5, 10, 20 clients", "Phần 4 — Chương Võ (TV4)")
    # Thêm bảng kết quả pgbench
    table_shape12 = s12.shapes.add_table(4, 6, PInches(0.8), PInches(1.8), PInches(11.7), PInches(2.5))
    t12 = table_shape12.table
    h12 = ["Số clients (-c)", "Tổng giao dịch", "Số GD thành công", "Tỷ lệ lỗi / Deadlock", "Độ trễ TB (Latency)", "Thông lượng (TPS)"]
    for j, h in enumerate(h12):
        cell = t12.cell(0, j)
        cell.text = h
        cell.fill.solid()
        cell.fill.fore_color.rgb = C_NAVY
        for p in cell.text_frame.paragraphs:
            p.font.bold = True
            p.font.size = PPt(12)
            p.font.color.rgb = PRGBColor(255, 255, 255)
    
    rows12 = [
        ["5 clients", "500", "500 / 500 (100%)", "0.00%", "1.448 ms", "3.452,90 TPS"],
        ["10 clients", "1.000", "1.000 / 1.000 (100%)", "0.00%", "2.912 ms", "3.434,18 TPS"],
        ["20 clients", "2.000", "2.000 / 2.000 (100%)", "0.00%", "5.814 ms", "3.439,94 TPS"],
    ]
    for i, row in enumerate(rows12):
        for j, val in enumerate(row):
            cell = t12.cell(i+1, j)
            cell.text = val
            for p in cell.text_frame.paragraphs:
                p.font.size = PPt(12)
                p.font.color.rgb = C_TEXT_DARK

    tb12_sub = s12.shapes.add_textbox(PInches(0.8), PInches(4.6), PInches(11.7), PInches(2.2))
    tf12_sub = tb12_sub.text_frame
    tf12_sub.word_wrap = True
    tf12_sub.text = "• Nhận xét kết quả thực nghiệm:\n" \
                    "   1. Thông lượng cực kỳ ổn định ở mức ~3.440 TPS ngay cả khi số lượng client tăng gấp 4 lần.\n" \
                    "   2. Độ trễ tăng tuyến tính từ 1.44ms lên 5.81ms, hoàn toàn đáp ứng tiêu chuẩn thời gian thực (< 10ms).\n" \
                    "   3. Tỷ lệ Deadlock = 0% nhờ quy ước chuẩn hóa thứ tự khóa user_a < user_b trong proc_add_friend."
    for p in tf12_sub.paragraphs:
        p.font.size = PPt(13)
        p.font.color.rgb = C_TEXT_DARK

    # SLIDE 13: THIẾT KẾ VẬT LÝ & KẾ HOẠCH ĐÁNH CHỈ MỤC (TV5 - QUỲNH)
    s13 = prs.slides.add_slide(blank_layout)
    add_header(s13, "Thiết Kế Vật Lý & Kế Hoạch Đánh Chỉ Mục (Indexing)", 
               "Tối ưu hóa truy vấn cho từng mô hình và đo lường dung lượng đĩa bổ sung", "Phần 5 — Như Quỳnh (TV5)")
    tb13 = s13.shapes.add_textbox(PInches(0.8), PInches(1.8), PInches(11.7), PInches(5.0))
    tf13 = tb13.text_frame
    tf13.word_wrap = True
    tf13.text = "• Chiến lược đánh chỉ mục theo đặc thù từng mô hình (indexes.sql):\n" \
                "   - Adjacency List: B-tree trên comment(parent_id) và (post_id, parent_id) => Tăng tốc độ lọc điểm neo và phép nối CTE.\n" \
                "   - Path Enumeration: GiST index trên cột kiểu ltree comment_path(path) => Quét mẫu tiền tố đường dẫn siêu tốc.\n" \
                "   - Nested Set: Composite Index trên comment_ns(post_id, lft, rgt) => Chuyển Seq Scan sang Index Only Scan khi lọc BETWEEN.\n" \
                "   - Closure Table: Hai chỉ mục idx_closure_ancestor_depth và idx_closure_descendant_ancestor.\n\n" \
                "• Đo lường dung lượng đĩa và chỉ mục qua storage_sizes.sql (pg_relation_size):\n" \
                "   - Ở quy mô 100.000 bình luận: Bảng comment chỉ chiếm 100.000 dòng; trong khi comment_closure chiếm 1.067.146 dòng (gấp 10.6 lần) do phải lưu toàn bộ các cặp tổ tiên - hậu duệ.\n" \
                "   - Dung lượng chỉ mục của Closure Table lớn nhất; Adjacency List tiêu tốn ít dung lượng đĩa nhất."
    for p in tf13.paragraphs:
        p.font.size = PPt(13)
        p.font.color.rgb = C_TEXT_DARK

    # SLIDE 14: KẾT QUẢ BENCHMARK 4 MÔ HÌNH CÂY (TV5 - QUỲNH)
    s14 = prs.slides.add_slide(blank_layout)
    add_header(s14, "Kết Quả Đo Đạc Benchmark 4 Mô Hình Cây Phân Cấp", 
               "Tổng hợp 360 phép đo thực nghiệm EXPLAIN (ANALYZE, BUFFERS) trên 1k, 10k, 100k nút", "Phần 5 — Như Quỳnh (TV5)")
    
    # Text bên trái, biểu đồ bên phải
    tb14 = s14.shapes.add_textbox(PInches(0.8), PInches(1.8), PInches(6.2), PInches(5.0))
    tf14 = tb14.text_frame
    tf14.word_wrap = True
    tf14.text = "• Thời gian đọc cây con Q2 (Subtree):\n" \
                "   - Khi có chỉ mục: Nested Set (~0.05ms) và Closure Table (~0.04ms) đọc nhanh nhất.\n" \
                "   - Path Enumeration (ltree) đạt ~0.15ms nhờ GiST index.\n" \
                "   - Adjacency List chậm hơn (~1.2ms – 3.5ms) do chi phí vòng lặp CTE đệ quy tăng theo độ sâu.\n\n" \
                "• Chi phí ghi (INSERT và MOVE SUBTREE):\n" \
                "   - Adjacency List ghi nhanh nhất (chỉ ghi 1 dòng O(1)).\n" \
                "   - Nested Set và Closure Table ghi chậm nhất do phải re-index hoặc nhân bản nhiều dòng bao đóng."
    for p in tf14.paragraphs:
        p.font.size = PPt(12.5)
        p.font.color.rgb = C_TEXT_DARK

    img_bench = ROOT / "benchmark" / "results" / "charts" / "q2_subtree_10000.png"
    if img_bench.exists():
        s14.shapes.add_picture(str(img_bench), PInches(7.2), PInches(1.8), width=PInches(5.3))

    # SLIDE 15: MA TRẬN ĐÁNH ĐỔI (TRADE-OFF) & KHUYẾN NGHỊ KIẾN TRÚC (TV5 + CẢ NHÓM)
    s15 = prs.slides.add_slide(blank_layout)
    add_header(s15, "Ma Trận Đánh Đổi (Trade-Off) & Khuyến Nghị Kiến Trúc", 
               "Lựa chọn công nghệ tối ưu dựa trên bài toán thực tế của Mạng xã hội học tập", "Tổng kết — Cả nhóm")
    tb15 = s15.shapes.add_textbox(PInches(0.8), PInches(1.8), PInches(11.7), PInches(5.0))
    tf15 = tb15.text_frame
    tf15.word_wrap = True
    tf15.text = "• Ma trận đánh đổi tổng thể (Trade-off Matrix):\n" \
                "   [Tốc độ đọc cây con]:  Closure Table > Nested Set > Path Enumeration (ltree) > Adjacency List\n" \
                "   [Tốc độ ghi / cập nhật]: Adjacency List > Path Enumeration > Closure Table > Nested Set\n" \
                "   [Dung lượng đĩa tiêu tốn]: Adjacency List (nhẹ nhất) < Path Enumeration < Nested Set < Closure Table (nặng nhất)\n\n" \
                "• Khuyến nghị kiến trúc cho Mạng xã hội học tập (Tỷ lệ 90% Đọc, 10% Ghi):\n" \
                "   1. Hệ thống Bình luận (Threaded Comments): Chọn Path Enumeration (ltree) — Cân bằng hoàn hảo giữa đọc nhanh và ghi nhẹ.\n" \
                "   2. Danh mục Thư mục học liệu (Cây tĩnh ít sửa): Chọn Nested Set hoặc Closure Table để tối ưu hóa tốc độ xem tài liệu.\n" \
                "   3. Mạng lưới quan hệ bạn bè: Dùng PostgreSQL WITH RECURSIVE cho bài toán bạn của bạn độ sâu k <= 2; tích hợp Neo4j khi cần thuật toán gợi ý tri thức đa tầng."
    for p in tf15.paragraphs:
        p.font.size = PPt(13)
        p.font.color.rgb = C_TEXT_DARK

    # SLIDE 16: KẾT LUẬN & ĐÓNG GÓP CỦA ĐỒ ÁN (TV1 - QUANG)
    s16 = prs.slides.add_slide(blank_layout)
    add_header(s16, "Kết Luận, Đóng Góp Của Đồ Án & Hướng Phát Triển", 
               "Đáp ứng toàn diện 100% chuẩn đầu ra môn học và khả năng tái lập kết quả độc lập", "Phần 1 — Trần Quang (Leader)")
    tb16 = s16.shapes.add_textbox(PInches(0.8), PInches(1.8), PInches(11.7), PInches(5.0))
    tf16 = tb16.text_frame
    tf16.word_wrap = True
    tf16.text = "• Kết quả đạt được:\n" \
                "   1. Xây dựng hoàn chỉnh CSDL Mạng xã hội học tập chuẩn BCNF và 4NF với đầy đủ cơ chế toàn vẹn Trigger/Procedure.\n" \
                "   2. Hiện thực và đối chứng thực nghiệm 4 mô hình cây phân cấp trên quy mô 1.000 đến 100.000 nút.\n" \
                "   3. Làm rõ ranh giới giữa SQL đệ quy PostgreSQL và Neo4j Cypher.\n" \
                "   4. Thực nghiệm thành công bế tắc Deadlock (40P01) và chứng minh hệ thống chịu tải 3.440 TPS qua pgbench.\n" \
                "   5. Đầy đủ minh chứng làm việc nhóm (CLO7): 8 biên bản họp tuần, lịch sử commit minh bạch, review chéo chặt chẽ.\n\n" \
                "• Hướng phát triển tiếp theo:\n" \
                "   - Nghiên cứu kiến trúc Đa mô hình lưu trữ (Polyglot Persistence): Kết hợp PostgreSQL (ACID) + Neo4j (Graph Recommender).\n" \
                "   - Áp dụng phân vùng bảng (Table Partitioning) cho bảng Closure Table khi dữ liệu vượt ngưỡng 10 triệu bản ghi."
    for p in tf16.paragraphs:
        p.font.size = PPt(13)
        p.font.color.rgb = C_TEXT_DARK

    # SLIDE 17: KỊCH BẢN DEMO THỰC TẾ (TV5 + TV3)
    s17 = prs.slides.add_slide(blank_layout)
    add_header(s17, "Kịch Bản Demo Thực Tế Hệ Thống (5 Phút)", 
               "Quy trình tái lập từ CSDL sạch và kiểm chứng trực tiếp trước Hội đồng", "Demo — TV5 & TV3")
    tb17 = s17.shapes.add_textbox(PInches(0.8), PInches(1.8), PInches(11.7), PInches(5.0))
    tf17 = tb17.text_frame
    tf17.word_wrap = True
    tf17.text = "• Bước 1 (30s): Khởi tạo CSDL sạch từ schema.sql, constraints.sql và seed_test.sql.\n\n" \
                "• Bước 2 (60s): Demo tính đồng nhất của 4 mô hình cây: Chạy câu Q2 (lấy cây con bình luận 1) trên cả 4 mô hình, chứng minh cùng trả về tập kết quả {1, 2, 3}.\n\n" \
                "• Bước 3 (60s): Demo SQL đệ quy: Chạy bài toán R2 tìm đường đi ngắn nhất An -> Em ra [1, 3, 5] và đối chiếu với câu lệnh Cypher trên Neo4j Browser.\n\n" \
                "• Bước 4 (60s): Kích hoạt Deadlock trực tiếp: Mở 2 cửa sổ psql chạy Session A và Session B đối kháng, kích hoạt mã lỗi 40P01 trước màn hình.\n\n" \
                "• Bước 5 (90s): Trình chiếu bộ công cụ Benchmark tự động: Chạy run_benchmark.py và hiển thị biểu đồ so sánh tự động sinh bằng matplotlib."
    for p in tf17.paragraphs:
        p.font.size = PPt(13)
        p.font.color.rgb = C_TEXT_DARK

    # SLIDE 18: HỎI ĐÁP PHẢN BIỆN (Q&A)
    s18 = prs.slides.add_slide(blank_layout)
    bg18 = s18.shapes.add_shape(MSO_SHAPE.RECTANGLE, 0, 0, PInches(13.333), PInches(7.5))
    bg18.fill.solid()
    bg18.fill.fore_color.rgb = C_NAVY
    bg18.line.color.rgb = C_NAVY

    tb18 = s18.shapes.add_textbox(PInches(1.5), PInches(2.2), PInches(10.3), PInches(3.5))
    tf18 = tb18.text_frame
    tf18.word_wrap = True
    
    p = tf18.paragraphs[0]
    p.text = "XIN CHÂN THÀNH CẢM ƠN QUÝ THẦY CÔ VÀ CÁC BẠN!"
    p.font.size = PPt(28)
    p.font.bold = True
    p.font.color.rgb = PRGBColor(255, 255, 255)
    p.alignment = PP_ALIGN.CENTER

    p2 = tf18.add_paragraph()
    p2.text = "\nNHÓM ĐỀ TÀI 17 SẴN SÀNG LẮNG NGHE Ý KIẾN ĐÓNG GÓP VÀ CÂU HỎI PHẢN BIỆN"
    p2.font.size = PPt(18)
    p2.font.color.rgb = PRGBColor(147, 197, 253)
    p2.alignment = PP_ALIGN.CENTER

    out_pptx = ROOT / "slides" / "De_tai_17_Slide_Thuyet_trinh.pptx"
    prs.save(str(out_pptx))
    print(f"-> Đã xuất Slide PowerPoint hoàn chỉnh: {out_pptx} ({os.path.getsize(out_pptx)} bytes)")

if __name__ == "__main__":
    build_docx_report()
    build_pptx_slides()
