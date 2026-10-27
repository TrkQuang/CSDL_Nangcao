"""Export project report Markdown to DOCX with headings, tables, code and images."""

from __future__ import annotations

import argparse
import re
from pathlib import Path

from docx import Document
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt


def inline(paragraph, text: str) -> None:
    tokens = re.split(r"(\*\*[^*]+\*\*|`[^`]+`|\[[^\]]+\]\([^)]+\))", text)
    for token in tokens:
        if token.startswith("**") and token.endswith("**"):
            paragraph.add_run(token[2:-2]).bold = True
        elif token.startswith("`") and token.endswith("`"):
            paragraph.add_run(token[1:-1]).font.name = "Consolas"
        else:
            link = re.fullmatch(r"\[([^\]]+)\]\(([^)]+)\)", token)
            paragraph.add_run(f"{link[1]} ({link[2]})" if link else token)


def export(source: Path, destination: Path) -> None:
    document = Document()
    section = document.sections[0]
    section.page_width, section.page_height = Cm(21), Cm(29.7)
    section.top_margin = section.bottom_margin = Cm(2)
    section.left_margin, section.right_margin = Cm(2.5), Cm(2)
    normal = document.styles["Normal"]
    normal.font.name, normal.font.size = "Times New Roman", Pt(12)
    normal.paragraph_format.space_after = Pt(6)
    normal.paragraph_format.line_spacing = 1.15
    for name in ("Title", "Heading 1", "Heading 2", "Heading 3"):
        document.styles[name].font.name = "Times New Roman"
    footer = section.footer.paragraphs[0]
    footer.alignment = 2
    footer.add_run("Trang ")
    field = OxmlElement("w:fldSimple")
    field.set(qn("w:instr"), "PAGE")
    footer._p.append(field)
    lines = source.read_text(encoding="utf-8").splitlines()
    position = 0
    code = False
    while position < len(lines):
        line = lines[position]
        position += 1
        if line.startswith("```"):
            code = not code
            continue
        if code:
            paragraph = document.add_paragraph()
            paragraph.paragraph_format.space_after = Pt(0)
            run = paragraph.add_run(line)
            run.font.name, run.font.size = "Consolas", Pt(8)
            continue
        if not line.strip() or line.strip() == "---":
            continue
        heading = re.match(r"^(#{1,4}) (.+)$", line)
        image = re.fullmatch(r"!\[([^\]]*)\]\(([^)]+)\)", line)
        if heading:
            level = len(heading[1]) - 1
            document.add_heading(heading[2], level=level)
        elif image:
            image_path = (source.parent / image[2]).resolve()
            if not image_path.is_file():
                raise FileNotFoundError(image_path)
            document.add_picture(str(image_path), width=Cm(16))
            document.add_paragraph(image[1], style="Caption")
        elif line.startswith("|"):
            rows = [line]
            while position < len(lines) and lines[position].startswith("|"):
                rows.append(lines[position])
                position += 1
            cells = [[cell.strip() for cell in row.strip().strip("|").split("|")] for row in rows]
            cells = [row for row in cells if not all(re.fullmatch(r":?-+:?", cell) for cell in row)]
            if not cells:
                continue
            table = document.add_table(rows=0, cols=len(cells[0]))
            table.style = "Table Grid"
            for row_index, row in enumerate(cells):
                if len(row) != len(cells[0]):
                    raise ValueError("Inconsistent Markdown table width")
                target = table.add_row().cells
                for cell, value in zip(target, row):
                    inline(cell.paragraphs[0], value)
                    for run in cell.paragraphs[0].runs:
                        run.font.size = Pt(9)
                        if row_index == 0:
                            run.bold = True
                if row_index == 0:
                    repeat = OxmlElement("w:tblHeader")
                    table.rows[0]._tr.get_or_add_trPr().append(repeat)
            document.add_paragraph()
        elif line.startswith("- "):
            inline(document.add_paragraph(style="List Bullet"), line[2:])
        elif re.match(r"^\d+\. ", line):
            inline(document.add_paragraph(style="List Number"), re.sub(r"^\d+\. ", "", line))
        else:
            inline(document.add_paragraph(), line.lstrip("> "))
    if code:
        raise ValueError("Unclosed Markdown code fence")
    document.core_properties.title = next((line[2:] for line in lines if line.startswith("# ")), source.stem)
    document.core_properties.author = "Quỳnh"
    document.core_properties.subject = "Đề tài 17 - Dữ liệu, thiết kế vật lý và benchmark"
    destination.parent.mkdir(parents=True, exist_ok=True)
    document.save(destination)
    reopened = Document(destination)
    if not reopened.paragraphs:
        raise ValueError("Empty DOCX export")
    print(f"Exported {destination}: {len(reopened.paragraphs)} paragraphs, {len(reopened.tables)} tables, {len(reopened.inline_shapes)} images")


def main() -> None:
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("source", type=Path)
    parser.add_argument("destination", type=Path)
    args = parser.parse_args()
    export(args.source.resolve(), args.destination.resolve())


if __name__ == "__main__":
    main()
